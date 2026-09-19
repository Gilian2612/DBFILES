import os
import pdfplumber
from docx import Document
import openpyxl
from pptx import Presentation

def extraer_texto(ruta: str, tipo: str) -> str:
    """Extrae el texto plano de un archivo según su tipo."""
    try:
        if tipo == "pdf":
            return _parse_pdf(ruta)
        elif tipo == "docx":
            return _parse_docx(ruta)
        elif tipo == "xlsx":
            return _parse_xlsx(ruta)
        elif tipo == "pptx":
            return _parse_pptx(ruta)
        else:
            return ""
    except Exception as e:
        print(f"[Parser] Error procesando {ruta}: {e}")
        return ""

def _parse_pdf(ruta: str) -> str:
    textos = []
    with pdfplumber.open(ruta) as pdf:
        for pagina in pdf.pages:
            texto = pagina.extract_text()
            if texto:
                textos.append(texto)
    return "\n".join(textos)

def _parse_docx(ruta: str) -> str:
    doc = Document(ruta)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def _parse_xlsx(ruta: str) -> str:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    textos = []
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            celdas = [str(c) for c in fila if c is not None and str(c).strip()]
            if celdas:
                textos.append(" ".join(celdas))
    return "\n".join(textos)

def _parse_pptx(ruta: str) -> str:
    prs = Presentation(ruta)
    textos = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                textos.append(shape.text)
    return "\n".join(textos)

def detectar_tipo(nombre_archivo: str) -> str:
    ext = os.path.splitext(nombre_archivo)[1].lower()
    mapa = {
        ".pdf":  "pdf",
        ".docx": "docx",
        ".doc":  "docx",
        ".xlsx": "xlsx",
        ".xls":  "xlsx",
        ".pptx": "pptx",
        ".ppt":  "pptx",
    }
    return mapa.get(ext, "otro")
